from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from PIL import Image as PILImage

from seller_auto.workbook import (
    TEMPLATE_SHEET,
    can_enter_top,
    dimension_cm,
    fill_package_rows,
    lowest_selling_price,
    ranked_matches,
    sales_cutoff,
    template_path,
    weight_kg,
    write_product_workbook,
)


def _column(sheet, header_prefix: str, occurrence: int = 1) -> int:
    seen = 0
    for cell in next(sheet.iter_rows(min_row=1, max_row=1)):
        text = "".join(str(cell.value or "").split())
        if text.startswith(header_prefix):
            seen += 1
            if seen == occurrence:
                return cell.column
    raise AssertionError(header_prefix)


def _mini_workbook(path: Path, existing: str | None = "已有产品") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TEMPLATE_SHEET
    headers = [
        "名称",
        "SKU",
        "上架",
        "毛重(KG)",
        "长（cm）",
        "宽（cm）",
        "高（cm）",
        "体积（CBM)",
        "头程运输方式",
        "汇率",
        "售价（MX$)",
        "渠道比例   (%)",
        "采购成本   (RMB)",
        "渠道费用",
        "头程运费(RMB)",
        "尾程总运费(RMB)",
        "FBA\n运费",
        "FBM运费（$）",
        "调拨费",
        "仓库处理费(RMB)",
        "税20%    (RMB)",
        "提现汇损9%",
        "代扣代缴19%",
        "销售价格(RMB)",
        "成本合计(RMB)",
        "毛利润",
        "毛利率",
        "投产比",
        "MKD链接",
        "MKD链接",
        "MKD链接",
        "预计数量",
        "预计方数",
        "预计货值",
    ]
    for index, header in enumerate(headers, start=1):
        sheet.cell(1, index, header)
    if existing:
        sheet.cell(2, 1, existing)
        sheet.cell(2, 2, "old-sku")
        sheet.cell(2, 32, 10)
        sheet.cell(2, 33, "=AF2*H2")
        sheet.cell(2, 34, "=AF2*M2")
    workbook.save(path)


def _product(image_path: str = "") -> dict:
    return {
        "chinese_name": "硅胶灶台缝隙盖",
        "keywords": "cubiertas hendidura para estufa",
        "image_path": image_path,
        "package_size": {
            "length": "11.2",
            "width": "9.7",
            "height": "6.1",
            "unit": "cm",
            "weight": "204",
            "weight_unit": "g",
        },
        "source_price": {"average": "12.35"},
        "matches": [
            {
                "candidate_id": "low-sales",
                "url": "https://www.mercadolibre.com.ar/p/MLA1?tracking=1#x",
                "sales_text": "1 vendido",
                "price": 99000,
                "confidence": 0.99,
            },
            {
                "candidate_id": "mid-sales",
                "url": "https://www.mercadolibre.com.ar/p/MLA2",
                "sales_text": "+50 vendidos",
                "price": 12000,
                "confidence": 0.9,
            },
            {
                "candidate_id": "top-sales",
                "url": "https://www.mercadolibre.com.ar/p/MLA3",
                "sales_text": "+500 vendidos",
                "price": 18000,
                "confidence": 0.88,
            },
        ],
    }


class WorkbookTests(unittest.TestCase):
    def test_weight_and_size_use_the_sheet_units(self) -> None:
        package = {
            "weight": "454",
            "weight_unit": "g",
            "length": "10",
            "width": "8",
            "height": "6",
            "unit": "in",
        }
        self.assertEqual(weight_kg(package), 0.454)
        self.assertEqual(dimension_cm(package, "length"), 25.4)

    def test_top_ten_are_sorted_by_sales_and_use_the_lowest_price(self) -> None:
        top = ranked_matches(_product())
        self.assertEqual([item["candidate_id"] for item in top], ["top-sales", "mid-sales", "low-sales"])
        self.assertEqual(lowest_selling_price(top), 12000)

    def test_sales_cutoff_waits_until_ten_matches(self) -> None:
        product = _product()
        self.assertIsNone(sales_cutoff(product))
        self.assertTrue(can_enter_top(product, "1 vendido"))
        product["matches"] = [
            {
                "candidate_id": f"m{index}",
                "url": f"https://www.mercadolibre.com.ar/p/MLA{index}",
                "key": f"MLA{index}",
                "sales_text": "+100 vendidos",
                "source": "search",
            }
            for index in range(10)
        ]
        self.assertEqual(sales_cutoff(product), 100)
        self.assertFalse(can_enter_top(product, "+50 vendidos"))
        self.assertTrue(can_enter_top(product, "+100 vendidos"))
        self.assertTrue(can_enter_top(product, "+500 vendidos"))

    def test_appends_a_row_with_price_links_and_listing_image(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            image_path = Path(directory) / "hero.jpg"
            PILImage.new("RGB", (40, 40), (20, 20, 20)).save(image_path, "JPEG")
            _mini_workbook(path)
            write_product_workbook(path, [_product(str(image_path))])
            sheet = load_workbook(path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(2, 1).value, "已有产品")
            self.assertEqual(sheet.cell(3, _column(sheet, "名称")).value, "硅胶灶台缝隙盖")
            self.assertEqual(sheet.cell(3, _column(sheet, "SKU")).value, "cubiertas hendidura para estufa")
            self.assertEqual(sheet.cell(3, _column(sheet, "毛重")).value, 0.204)
            self.assertEqual(sheet.cell(3, _column(sheet, "长")).value, 11.2)
            self.assertEqual(sheet.cell(3, _column(sheet, "采购成本")).value, 12.35)
            self.assertEqual(sheet.cell(3, _column(sheet, "头程运输方式")).value, "AR普货海运")
            self.assertEqual(sheet.cell(3, _column(sheet, "售价")).value, 12000)
            self.assertEqual(sheet.cell(3, _column(sheet, "汇率")).value, 0.004)
            self.assertEqual(sheet.cell(3, _column(sheet, "销售价格")).value, "=K3*J3")
            mkd_headers = [
                cell.value
                for cell in next(sheet.iter_rows(min_row=1, max_row=1))
                if "".join(str(cell.value or "").split()) == "MKD链接"
            ]
            self.assertEqual(len(mkd_headers), 10)
            self.assertEqual(sheet.cell(3, _column(sheet, "MKD链接", 1)).value, "https://www.mercadolibre.com.ar/p/MLA3")
            self.assertEqual(sheet.cell(3, _column(sheet, "MKD链接", 2)).value, "https://www.mercadolibre.com.ar/p/MLA2")
            self.assertEqual(sheet.cell(3, _column(sheet, "MKD链接", 3)).value, "https://www.mercadolibre.com.ar/p/MLA1")
            self.assertEqual(len(sheet._images), 1)
            self.assertEqual(sheet._images[0].anchor._from.col, 2)
            self.assertEqual(sheet._images[0].anchor._from.row, 2)
            self.assertEqual(sheet.cell(2, _column(sheet, "预计方数")).value, "=AM2*H2")
            self.assertEqual(sheet.cell(2, _column(sheet, "预计货值")).value, "=AM2*M2")
            self.assertIsNone(sheet.cell(3, _column(sheet, "FBA")).value)

    def test_lingdong_orders_and_last_mile_fill_name_and_fba(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            _mini_workbook(path)
            product = _product()
            product["chinese_name"] = "硅胶商品"
            product["market_orders"] = 50
            product["fba_fee"] = "2740.00"
            write_product_workbook(path, [product])
            sheet = load_workbook(path)[TEMPLATE_SHEET]
            name = sheet.cell(3, _column(sheet, "名称"))
            self.assertEqual(name.value, "硅胶商品\n（50单）")
            self.assertTrue(name.alignment.wrap_text)
            self.assertEqual(name.alignment.horizontal, "center")
            self.assertEqual(name.alignment.vertical, "center")
            self.assertEqual(sheet.cell(3, _column(sheet, "SKU")).alignment.horizontal, "center")
            self.assertEqual(sheet.cell(3, _column(sheet, "采购成本")).alignment.vertical, "center")
            self.assertEqual(sheet.cell(3, _column(sheet, "预计数量")).value, 50)
            self.assertEqual(sheet.cell(3, _column(sheet, "预计方数")).value, "=AM3*H3")
            self.assertEqual(sheet.cell(3, _column(sheet, "预计货值")).value, "=AM3*M3")
            self.assertEqual(sheet.cell(3, _column(sheet, "FBA")).value, 2740)
            self.assertTrue(str(sheet.cell(3, _column(sheet, "尾程")).value).startswith("="))

    def test_cost_and_fba_keep_the_previous_fill_and_roi_stays_clear(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            _mini_workbook(path)
            workbook = load_workbook(path)
            sheet = workbook[TEMPLATE_SHEET]
            yellow = PatternFill(fill_type="solid", fgColor="FFFF00")
            sheet.cell(2, _column(sheet, "采购成本")).fill = yellow
            sheet.cell(2, _column(sheet, "FBA")).fill = yellow
            sheet.cell(2, _column(sheet, "投产比")).fill = PatternFill(fill_type="solid", fgColor="00B0F0")
            workbook.save(path)
            product = _product()
            product["fba_fee"] = 2740
            write_product_workbook(path, [product])
            sheet = load_workbook(path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(3, _column(sheet, "采购成本")).fill.fgColor.rgb, "00FFFF00")
            self.assertEqual(sheet.cell(3, _column(sheet, "FBA")).fill.fgColor.rgb, "00FFFF00")
            self.assertIsNone(sheet.cell(3, _column(sheet, "投产比")).fill.patternType)
            self.assertEqual(sheet.cell(3, _column(sheet, "投产比")).number_format, "0.00%")

    def test_formula_columns_copy_the_previous_row(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            _mini_workbook(path)
            workbook = load_workbook(path)
            sheet = workbook[TEMPLATE_SHEET]
            sheet.cell(2, _column(sheet, "代扣代缴"), "=X2*0.23")
            sheet.cell(2, _column(sheet, "尾程"), "=(Q2+R2+S2)*J2")
            workbook.save(path)
            product = _product()
            product["fba_fee"] = 2740
            write_product_workbook(path, [product])
            sheet = load_workbook(path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(3, _column(sheet, "代扣代缴")).value, "=X3*0.19")
            self.assertEqual(sheet.cell(3, _column(sheet, "尾程")).value, "=(Q3+R3+S3)*J3")
            self.assertEqual(sheet.cell(3, _column(sheet, "FBA")).value, 2740)
            self.assertEqual(sheet.cell(3, _column(sheet, "销售价格")).value, "=K3*J3")

    def test_products_without_matches_are_not_written(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            _mini_workbook(path)
            empty = _product()
            empty["chinese_name"] = "嵌入式地埋灯"
            empty["matches"] = []
            written = write_product_workbook(path, [empty, _product()])
            sheet = load_workbook(path)[TEMPLATE_SHEET]
            self.assertEqual(written, 1)
            self.assertEqual(sheet.cell(2, 1).value, "已有产品")
            self.assertEqual(sheet.cell(3, _column(sheet, "名称")).value, "硅胶灶台缝隙盖")
            self.assertIsNone(sheet.cell(4, 1).value)
            self.assertEqual(write_product_workbook(path, [empty]), 0)
            sheet = load_workbook(path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(3, _column(sheet, "名称")).value, "硅胶灶台缝隙盖")
            self.assertIsNone(sheet.cell(4, 1).value)

    def test_fill_package_rows_updates_the_matching_sku(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            _mini_workbook(path)
            empty = _product()
            empty["package_size"] = None
            empty["source_price"] = None
            write_product_workbook(path, [empty])
            filled = _product()
            filled["fba_fee"] = 1390
            filled["market_orders"] = 12
            self.assertEqual(fill_package_rows(path, [filled]), 1)
            sheet = load_workbook(path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(3, _column(sheet, "毛重")).value, 0.204)
            self.assertEqual(sheet.cell(3, _column(sheet, "长")).value, 11.2)
            self.assertEqual(sheet.cell(3, _column(sheet, "宽")).value, 9.7)
            self.assertEqual(sheet.cell(3, _column(sheet, "高")).value, 6.1)
            self.assertEqual(sheet.cell(3, _column(sheet, "采购成本")).value, 12.35)
            self.assertEqual(sheet.cell(3, _column(sheet, "FBA")).value, 1390)
            self.assertEqual(sheet.cell(3, _column(sheet, "名称")).value, "硅胶灶台缝隙盖\n（12单）")
            self.assertEqual(sheet.cell(2, 1).value, "已有产品")
            later = _product()
            later["market_orders"] = 99
            self.assertEqual(fill_package_rows(path, [later]), 0)
            sheet = load_workbook(path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(3, _column(sheet, "名称")).value, "硅胶灶台缝隙盖\n（12单）")

    def test_template_file_is_still_the_working_sheet(self) -> None:
        self.assertEqual(template_path().name, "test.xlsx")


if __name__ == "__main__":
    unittest.main()
