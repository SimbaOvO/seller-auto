from __future__ import annotations

import os
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from seller_auto.lingdong import (
    _GATE_JS,
    _TABLE_JS,
    _query_cooldown,
    apply_last_mile,
    chargeable_kilograms,
    choose_quote_row,
    item_id,
    last_mile_fee,
    link_batches,
    package_fields,
    parse_ars_amount,
    quote_links,
    table_decision,
    total_orders,
)
from seller_auto.pacing import Cooldown, DelaySettings
from seller_auto.service import SellerAutoService, _new_state
from seller_auto.workbook import TEMPLATE_SHEET, display_name


def _chrome_for_testing() -> str:
    from seller_auto.browser_launch import chrome_for_testing_binary, default_playwright_browsers_dir

    binary = chrome_for_testing_binary(default_playwright_browsers_dir())
    if binary is None:
        raise unittest.SkipTest("本机没有 Chrome for Testing")
    return str(binary)


def _wait_export(service: object) -> dict:
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        result = (getattr(service, "state", None) or {}).get("export_result") or {}
        if result.get("message"):
            return result
        time.sleep(0.05)
    raise AssertionError("导出没有在预期时间内写完")


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TEMPLATE_SHEET
    sheet.cell(1, 1, "名称")
    sheet.cell(1, 2, "FBA\n运费")
    workbook.save(path)


class LingdongParseTests(unittest.TestCase):
    def test_amounts_and_item_ids(self) -> None:
        self.assertEqual(parse_ars_amount("商品售价 < 23999 并且 >= 16000：ARS 2740"), 2740)
        self.assertEqual(parse_ars_amount("尾程运费 ¥11.98 ARS2,740.00 其他成本"), 2740)
        self.assertEqual(parse_ars_amount("ARS 0"), 0)
        self.assertIsNone(parse_ars_amount("请选择"))

    def test_links_are_split_into_batches(self) -> None:
        links = [f"https://www.mercadolibre.com.ar/p/MLA{index}" for index in range(7)]
        batches = link_batches(links, size=5)
        self.assertEqual([len(batch) for batch in batches], [5, 2])
        self.assertEqual(link_batches(["", "https://example.com/a"]), [["https://example.com/a"]])
        self.assertEqual(link_batches([]), [])

    def test_rate_table_matches_the_calculator(self) -> None:
        self.assertEqual(last_mile_fee(0.22, 22425), 2740)
        self.assertEqual(last_mile_fee(0, 118756.63), 8148)
        self.assertEqual(last_mile_fee(2.798, 118756.63), 11508)
        self.assertEqual(last_mile_fee(250, 118756.63), 147948)
        self.assertEqual(last_mile_fee(0.3, 22425), 2740)
        self.assertEqual(last_mile_fee(0.31, 22425), 2760)
        self.assertEqual(last_mile_fee(180, 60000), 135588)
        self.assertEqual(last_mile_fee(180.1, 60000), 147948)
        heavy = chargeable_kilograms(
            {"item_weight": "200", "item_length": "100 cm", "item_width": "100 cm", "item_height": "100 cm"},
            None,
        )
        self.assertEqual(heavy, 250)
        self.assertEqual(last_mile_fee(heavy, 118756.63), 147948)
        package = {"length": "10", "width": "10", "height": "10", "unit": "cm", "weight": "200", "weight_unit": "g"}
        self.assertEqual(chargeable_kilograms({}, package), 0.25)
        listing = {"item_weight": "5000", "item_length": "100 cm", "item_width": "100 cm", "item_height": "100 cm"}
        self.assertEqual(chargeable_kilograms(listing, package), 0.25)
        self.assertEqual(item_id("https://articulo.mercadolibre.com.ar/MLA-2841522984-item-_JM"), "MLA2841522984")
        self.assertEqual(item_id("MLA2841522984"), "MLA2841522984")

    def test_orders_are_summed_once_per_item_and_freight_uses_the_lowest_price(self) -> None:
        rows = [
            {"id": "MLA1111111111", "sale30": 30, "price": 22000, "url": "", "source_url": "https://example.com/a"},
            {"id": "MLA1111111111", "sale30": 30, "price": 22000, "url": "", "source_url": "https://example.com/a2"},
            {"id": "MLA2222222222", "sale30": "50", "price": 18000, "url": "", "source_url": "https://example.com/b"},
            {"id": "MLA3333333333", "sale30": None, "price": None, "url": ""},
            {"id": "MLA4444444444", "sale30": 0, "price": 0, "url": "", "source_url": "https://example.com/c"},
        ]
        self.assertEqual(total_orders(rows), 80)
        self.assertEqual(choose_quote_row(rows)["id"], "MLA2222222222")
        self.assertIsNone(total_orders([{"id": "MLA3333333333", "sale30": None}]))

    def test_name_and_package_use_the_sheet_units(self) -> None:
        self.assertEqual(display_name({"chinese_name": "硅胶灶台缝隙盖"}), "硅胶灶台缝隙盖")
        self.assertEqual(display_name({"chinese_name": "硅胶灶台缝隙盖", "market_orders": 50}), "硅胶灶台缝隙盖\n（50单）")
        self.assertEqual(display_name({"chinese_name": "硅胶商品", "market_orders": 0}), "硅胶商品\n（0单）")
        fields = package_fields(
            {"length": "21.4", "width": "7.8", "height": "0.6", "unit": "cm", "weight": "213", "weight_unit": "g"}
        )
        self.assertEqual(fields["length"], 21.4)
        self.assertEqual(fields["weight"], 0.213)

    def test_each_query_waits_like_an_amazon_click(self) -> None:
        clock = {"t": 0.0}
        slept: list[float] = []
        samples = iter([9.0, 14.0, 8.0, 11.0])

        def sleep(seconds: float) -> None:
            slept.append(seconds)
            clock["t"] += seconds

        cooldown = Cooldown(
            8,
            15,
            sleep=sleep,
            now=lambda: clock["t"],
            rng=lambda _low, _high: next(samples),
            label="蓝鲸",
        )
        browser = SimpleNamespace(lingdong_cooldown=cooldown)
        _query_cooldown(browser).acquire()
        clock["t"] += 2
        _query_cooldown(browser).acquire()
        self.assertEqual(slept, [9.0, 12.0])

        clock["t"] = 100.0
        _query_cooldown(browser).acquire()
        self.assertEqual(slept, [9.0, 12.0])

    def test_query_cooldown_uses_the_amazon_interval(self) -> None:
        browser = SimpleNamespace(settings=DelaySettings())
        cooldown = _query_cooldown(browser)
        self.assertEqual((cooldown.min_seconds, cooldown.max_seconds), (8, 15))
        self.assertEqual(cooldown.label, "蓝鲸")
        self.assertIs(browser.lingdong_cooldown, cooldown)
        self.assertIsNotNone(cooldown.ready_at)

    def test_table_waits_for_a_real_refresh(self) -> None:
        before = "MLA1:10:100:False"
        same = {"loading": False, "empty": False, "rows": [{"id": "MLA1", "sale30": 10, "price": 100}]}
        updated = {"loading": False, "empty": False, "rows": [{"id": "MLA2", "sale30": 4, "price": 80}]}
        self.assertEqual(table_decision(before, {"loading": True, "rows": same["rows"]}, False, 0), "wait")
        self.assertEqual(table_decision(before, same, False, 2), "wait")
        self.assertEqual(table_decision(before, same, True, 0.2), "wait")
        self.assertEqual(table_decision(before, same, True, 0.8), "row")
        self.assertEqual(table_decision(before, updated, False, 0), "row")
        self.assertEqual(table_decision(before, {"loading": False, "empty": True, "rows": []}, False, 0), "empty")

    def test_quote_links_follow_the_sheet_order(self) -> None:
        product = {
            "matches": [
                {
                    "candidate_id": "low",
                    "url": "https://www.mercadolibre.com.ar/p/MLA1111111111?tracking=1#x",
                    "sales_text": "1 vendido",
                    "confidence": 0.9,
                },
                {
                    "candidate_id": "high",
                    "url": "https://articulo.mercadolibre.com.ar/MLA-2222222222-item-_JM",
                    "sales_text": "+50 vendidos",
                    "confidence": 0.95,
                },
            ]
        }
        self.assertEqual(
            quote_links(product),
            [
                "https://articulo.mercadolibre.com.ar/MLA-2222222222-item-_JM",
                "https://www.mercadolibre.com.ar/p/MLA1111111111",
            ],
        )


class LingdongExportTests(unittest.TestCase):
    def test_export_writes_orders_and_fba_without_opening_a_browser(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "book.xlsx"
            _workbook(workbook_path)
            service = SellerAutoService(DelaySettings(), output_dir=str(root), workbook_path=str(workbook_path))
            state = _new_state("sess", "SELLER1", "https://www.amazon.com.mx/s?me=SELLER1")
            state["single_product"] = "B0TEST1234"
            state["product_order"] = ["B0TEST1234"]
            state["current_product_id"] = "B0TEST1234"
            state["products"]["B0TEST1234"] = {
                "product_id": "B0TEST1234",
                "chinese_name": "硅胶灶台缝隙盖",
                "keywords": "cubiertas",
                "listing": {"page1_done": True, "page2_done": True, "candidates": [], "next_index": 0},
                "judged_ids": [],
                "package_size": {
                    "length": "21.4",
                    "width": "7.8",
                    "height": "0.6",
                    "unit": "cm",
                    "weight": "213",
                    "weight_unit": "g",
                },
                "matches": [
                    {
                        "candidate_id": "low",
                        "url": "https://www.mercadolibre.com.ar/p/MLA1111111111?tracking=1#x",
                        "sales_text": "1 vendido",
                        "price": 22000,
                        "confidence": 0.9,
                    },
                    {
                        "candidate_id": "high",
                        "url": "https://articulo.mercadolibre.com.ar/MLA-2222222222-item-_JM",
                        "sales_text": "+50 vendidos",
                        "price": 18000,
                        "confidence": 0.95,
                    },
                ],
            }
            service.state = state
            service._save()
            seen: dict = {}

            def fake(jobs: list[dict]) -> str:
                seen["jobs"] = jobs
                service.state["products"]["B0TEST1234"]["market_orders"] = 80
                service.state["products"]["B0TEST1234"]["fba_fee"] = 2740
                return "蓝鲸选品已把 30 天销量写入名称，并把尾程运费写入 FBA 运费。"

            service._run_market_quotes = fake
            service.export_results()
            payload_data = _wait_export(service)
            self.assertTrue(payload_data["ok"])
            self.assertIn("30 天销量", payload_data["message"])
            self.assertEqual(seen["jobs"][0]["package"]["weight"], "213")
            self.assertEqual(
                seen["jobs"][0]["links"],
                [
                    "https://articulo.mercadolibre.com.ar/MLA-2222222222-item-_JM",
                    "https://www.mercadolibre.com.ar/p/MLA1111111111",
                ],
            )
            sheet = load_workbook(workbook_path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(2, 1).value, "硅胶灶台缝隙盖\n（80单）")
            self.assertEqual(sheet.cell(2, 2).value, 2740)

    def test_package_freight_is_recorded_when_lingdong_fails(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "book.xlsx"
            _workbook(workbook_path)
            service = SellerAutoService(DelaySettings(), output_dir=str(root), workbook_path=str(workbook_path))
            state = _new_state("sess", "SELLER1", "https://www.amazon.com.mx/s?me=SELLER1")
            state["single_product"] = "B0TEST1234"
            state["product_order"] = ["B0TEST1234"]
            state["current_product_id"] = "B0TEST1234"
            product = {
                "product_id": "B0TEST1234",
                "chinese_name": "烫衣板套",
                "keywords": "funda",
                "listing": {"page1_done": True, "page2_done": True, "candidates": [], "next_index": 0},
                "judged_ids": [],
                "package_size": {
                    "length": "33",
                    "width": "27.1",
                    "height": "8.7",
                    "unit": "cm",
                    "weight": "260",
                    "weight_unit": "g",
                },
                "matches": [
                    {
                        "candidate_id": "low",
                        "url": "https://www.mercadolibre.com.ar/p/MLA45199092",
                        "sales_text": "+500 vendidos",
                        "price": 12400,
                        "confidence": 0.95,
                    },
                    {
                        "candidate_id": "high",
                        "url": "https://www.mercadolibre.com.ar/p/MLA856699073",
                        "sales_text": "+1000 vendidos",
                        "price": 19989,
                        "confidence": 0.95,
                    },
                ],
            }
            state["products"]["B0TEST1234"] = product
            service.state = state
            service._save()
            self.assertEqual(apply_last_mile(product), 1430)

            def fake(jobs: list[dict]) -> str:
                raise RuntimeError(f"页面上没有阿根廷站点 {jobs}")

            service._run_market_quotes = fake
            service.export_results()
            payload_data = _wait_export(service)
            self.assertTrue(payload_data["ok"])
            self.assertIn("单量这次没写", payload_data["message"])
            self.assertIn("尾程运费已按包装和最低售价写入 FBA 运费", payload_data["message"])
            sheet = load_workbook(workbook_path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(2, 1).value, "烫衣板套")
            self.assertEqual(sheet.cell(2, 2).value, 1430)

    def test_export_without_links_skips_lingdong(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "book.xlsx"
            _workbook(workbook_path)
            service = SellerAutoService(DelaySettings(), output_dir=str(root), workbook_path=str(workbook_path))
            state = _new_state("sess", "SELLER1", "https://www.amazon.com.mx/s?me=SELLER1")
            state["product_order"] = ["B0TEST1234"]
            state["current_product_id"] = "B0TEST1234"
            state["products"]["B0TEST1234"] = {
                "product_id": "B0TEST1234",
                "chinese_name": "硅胶灶台缝隙盖",
                "keywords": "cubiertas",
                "listing": {"page1_done": True, "page2_done": True, "candidates": [], "next_index": 0},
                "judged_ids": [],
                "matches": [],
            }
            service.state = state
            service._save()

            def fake(jobs: list[dict]) -> str:
                raise AssertionError(jobs)

            service._run_market_quotes = fake
            service.export_results()
            payload_data = _wait_export(service)
            self.assertTrue(payload_data["ok"])
            self.assertNotIn("蓝鲸", payload_data["message"])
            self.assertIn("没有同款，没有写入表格", payload_data["message"])
            sheet = load_workbook(workbook_path)[TEMPLATE_SHEET]
            self.assertIsNone(sheet.cell(2, 1).value)


    def test_each_finished_asin_is_appended_once(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "book.xlsx"
            _workbook(workbook_path)
            service = SellerAutoService(DelaySettings(), output_dir=str(root), workbook_path=str(workbook_path))
            state = _new_state("sess", "SELLER1", "https://www.amazon.com.mx/s?me=SELLER1")
            state["current_product_id"] = "B0FIRST000"
            state["product_order"] = ["B0FIRST000", "B0SECOND00"]
            state["products"]["B0FIRST000"] = _asin("折叠购物车", "carrito", "https://www.mercadolibre.com.ar/p/MLA1")
            second = _asin("硅胶灶台缝隙盖", "cubiertas", "https://www.mercadolibre.com.ar/p/MLA2")
            second["listing"] = {"page1_done": False, "page2_done": False, "candidates": [], "next_index": 0}
            state["products"]["B0SECOND00"] = second
            service.state = state
            service._save()
            seen: list[list[str]] = []

            def fake(jobs: list[dict]) -> str:
                seen.append([job["product_id"] for job in jobs])
                for job in jobs:
                    service.state["products"][job["product_id"]]["market_orders"] = 10
                return "蓝鲸选品已把 30 天销量写入名称。"

            service._run_market_quotes = fake
            service.export_results()
            _wait_export(service)
            sheet = load_workbook(workbook_path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(2, 1).value, "折叠购物车\n（10单）")
            self.assertIsNone(sheet.cell(3, 1).value)
            self.assertEqual(seen, [["B0FIRST000"]])
            self.assertTrue(service.state["products"]["B0FIRST000"]["workbook_written"])
            self.assertFalse(service.state["products"]["B0SECOND00"].get("workbook_written"))

            second["listing"] = {"page1_done": True, "page2_done": True, "candidates": [], "next_index": 0}
            service.state["current_product_id"] = "B0SECOND00"
            service._save()
            service.export_results()
            _wait_export(service)
            sheet = load_workbook(workbook_path)[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(2, 1).value, "折叠购物车\n（10单）")
            self.assertEqual(sheet.cell(3, 1).value, "硅胶灶台缝隙盖\n（10单）")
            self.assertEqual(seen[1], ["B0SECOND00"])

            service.export_results()
            sheet = load_workbook(workbook_path)[TEMPLATE_SHEET]
            self.assertIsNone(sheet.cell(4, 1).value)
            self.assertEqual(len(seen), 2)


def _asin(name: str, keywords: str, url: str) -> dict:
    return {
        "chinese_name": name,
        "keywords": keywords,
        "listing": {"page1_done": True, "page2_done": True, "candidates": [], "next_index": 0},
        "judged_ids": [],
        "matches": [
            {
                "candidate_id": "match",
                "url": url,
                "sales_text": "+10 vendidos",
                "price": 18000,
                "confidence": 0.95,
            }
        ],
    }


class LingdongPageTests(unittest.TestCase):
    def test_page_script_reads_sale_freight_and_the_link_box(self) -> None:
        if os.environ.get("SELLER_AUTO_BROWSER_TESTS") != "1":
            self.skipTest("默认不启动 Chrome for Testing")
        from playwright.sync_api import Error as PlaywrightError
        from playwright.sync_api import sync_playwright

        page_html = """
        <div class="flexRow pm16">
          <div class="flexRow flex1">
            <span class="label">商品ID</span>
            <input id="id-box" value="id">
          </div>
          <div class="flexRow flex1">
            <span class="label">商品链接</span>
            <div class="flex1"><input id="link-box" value=""></div>
          </div>
        </div>
        <div class="el-row">
          <button>重置</button>
          <button>查询</button>
        </div>
        <div class="vxe-table" id="grid"></div>
        """
        with sync_playwright() as playwright:
            try:
                browser = playwright.chromium.launch(executable_path=_chrome_for_testing(), headless=True)
            except PlaywrightError as exc:
                message = str(exc)
                if "Executable doesn't exist" in message or "has been closed" in message:
                    self.skipTest("当前环境无法启动 Chrome for Testing")
                raise
            try:
                page = browser.new_page()
                page.set_content(page_html)
                self.assertEqual(page.evaluate(_GATE_JS), "ready")
                page.evaluate(
                    """() => {
                      document.querySelector('#grid').__vueParentComponent = {
                        proxy: {
                          getTableData() {
                            return {tableData: [{id: 'MLA2841522984', sale30: 50, price: 22425, url: 'https://articulo.mercadolibre.com.ar/MLA-2841522984-item'}]};
                          }
                        }
                      };
                    }"""
                )
                table = page.evaluate(_TABLE_JS)
                self.assertEqual(table["rows"][0]["sale30"], 50)
                link = page.locator("span.label", has_text="商品链接").locator(
                    "xpath=ancestor::div[contains(@class,'flexRow')][1]//input"
                )
                self.assertEqual(link.get_attribute("id"), "link-box")
                page.evaluate("() => { location.hash = '#/login'; document.body.innerHTML = '<input type=\"password\">'; }")
                self.assertEqual(page.evaluate(_GATE_JS), "login")
                page.evaluate("() => { location.hash = ''; document.body.innerHTML = '<span class=\"label\">商品链接</span>'; }")
                self.assertIsNone(page.evaluate(_GATE_JS))
            finally:
                browser.close()


if __name__ == "__main__":
    unittest.main()
