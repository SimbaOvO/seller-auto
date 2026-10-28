import threading
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

from seller_auto.pacing import DelaySettings
from seller_auto.service import SellerAutoService, _new_state
from openpyxl import Workbook

from seller_auto.workbook import TEMPLATE_SHEET


def _product_state() -> dict:
    state = _new_state("sess", "A1A10HFD4KEZW5", "https://www.amazon.com.mx/s?me=A1A10HFD4KEZW5")
    state["single_product"] = "B07LCLQ47D"
    state["run_id"] = "run-1"
    state["product_order"] = ["B07LCLQ47D"]
    state["products"]["B07LCLQ47D"] = {
        "product_id": "B07LCLQ47D",
        "url": "https://www.amazon.com.mx/dp/B07LCLQ47D",
        "title": "Cubierta de Hendidura de Estufa de Silicona",
        "brand": "Yolispa",
        "keywords": "cubiertas hendidura para estufa",
        "chinese_name": "硅胶灶台缝隙盖",
        "package_size": None,
        "source_price": None,
        "listing": {"page1_done": True, "page2_done": True, "candidates": [], "next_index": 0},
        "matches": [],
    }
    return state


class EnrichmentTests(unittest.TestCase):
    def test_schedule_returns_before_lookup_finishes(self) -> None:
        release = threading.Event()
        with tempfile.TemporaryDirectory() as directory:
            service = SellerAutoService(DelaySettings(), output_dir=directory)
            service.state = _product_state()

            def fake_run(product_id: str, image_url: str, event: threading.Event) -> None:
                release.wait(timeout=2)
                event.set()

            service._run_enrichment = fake_run
            started = time.monotonic()
            service._schedule_enrichment("B07LCLQ47D", "https://example.com/img.jpg")
            self.assertLess(time.monotonic() - started, 0.2)
            release.set()
            service._wait_enrichment()

    def test_export_waits_for_background_sellersprite(self) -> None:
        started = threading.Event()
        release = threading.Event()
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "test.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = TEMPLATE_SHEET
            sheet.cell(1, 1, "名称")
            sheet.cell(1, 2, "SKU")
            workbook.save(workbook_path)
            service = SellerAutoService(DelaySettings(), output_dir=directory, workbook_path=str(workbook_path))
            service.state = _product_state()
            service._save()

            def fake_run(product_id: str, image_url: str, event: threading.Event) -> None:
                started.set()
                self.assertTrue(release.wait(timeout=2))
                with service._state_lock:
                    service.state["products"][product_id]["package_size"] = {
                        "length": "23.6",
                        "width": "16.3",
                        "height": "7.1",
                        "unit": "cm",
                        "weight": "454",
                        "weight_unit": "g",
                    }
                    service.state["products"][product_id]["source_price"] = {"average": "12.35"}
                    service._write_state()
                event.set()

            service._run_enrichment = fake_run
            service._schedule_enrichment("B07LCLQ47D", "https://example.com/img.jpg")
            self.assertTrue(started.wait(timeout=2))

            payload = service.export_results()
            self.assertEqual(payload.data["status"], "exporting")
            release.set()
            deadline = time.monotonic() + 5
            while time.monotonic() < deadline and not (service.state.get("export_result") or {}).get("ok"):
                time.sleep(0.05)
            result = service.state.get("export_result") or {}
            self.assertTrue(result.get("ok"))
            self.assertTrue(Path(result["workbook"]).is_file())
            self.assertEqual(service.state["products"]["B07LCLQ47D"]["package_size"]["weight"], "454")

    def test_source_price_comes_from_the_sellersprite_row(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            service = SellerAutoService(DelaySettings(), output_dir=directory)
            service.state = _product_state()

            def fake_open(browser: object, asin: str) -> tuple[dict[str, str], str]:
                self.assertEqual(asin, "B07LCLQ47D")
                return {"weight": "260", "weight_unit": "g"}, "lookup-page"

            def fake_source(browser: object, page: object, asin: str) -> dict[str, str]:
                self.assertEqual(page, "lookup-page")
                self.assertEqual(asin, "B07LCLQ47D")
                return {"average": "8.00", "currency": "CNY", "count": "1"}

            service._ensure_enrich_browser = lambda: object()
            service._open_competitor = fake_open
            service._read_1688_from_row = fake_source
            event = threading.Event()
            service._run_enrichment("B07LCLQ47D", "https://example.com/img.jpg", event)
            self.assertTrue(event.is_set())
            product = service.state["products"]["B07LCLQ47D"]
            self.assertEqual(product["package_size"]["weight"], "260")
            self.assertEqual(product["source_price"]["average"], "8.00")

    def test_closed_browser_is_reopened_and_queried_again(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            service = SellerAutoService(DelaySettings(), output_dir=directory)
            service.state = _product_state()
            calls = {"open": 0}

            def fake_open(browser: object, asin: str) -> tuple[dict[str, str], str]:
                calls["open"] += 1
                if calls["open"] == 1:
                    raise RuntimeError("BrowserContext.add_cookies: Target page, context or browser has been closed")
                return {"weight": "260", "weight_unit": "g", "length": "10", "width": "8", "height": "6", "unit": "cm"}, "lookup"

            service._ensure_enrich_browser = lambda: object()
            service._reset_enrich_browser = lambda: None
            service._open_competitor = fake_open
            service._read_1688_from_row = lambda browser, page, asin: {"average": "8.00"}
            event = threading.Event()
            service._run_enrichment("B07LCLQ47D", "https://example.com/img.jpg", event)
            self.assertTrue(event.is_set())
            self.assertEqual(calls["open"], 2)
            product = service.state["products"]["B07LCLQ47D"]
            self.assertEqual(product["package_size"]["length"], "10")
            self.assertEqual(product["source_price"]["average"], "8.00")

    def test_closed_browser_during_source_lookup_is_retried(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            service = SellerAutoService(DelaySettings(), output_dir=directory)
            service.state = _product_state()
            calls = {"open": 0, "source": 0}

            def fake_open(browser: object, asin: str) -> tuple[dict[str, str], str]:
                calls["open"] += 1
                return {"weight": "260", "weight_unit": "g", "length": "10", "width": "8", "height": "6", "unit": "cm"}, "lookup"

            def fake_source(browser: object, page: object, asin: str) -> dict[str, str]:
                calls["source"] += 1
                if calls["source"] == 1:
                    raise RuntimeError("BrowserContext.new_page: Target page, context or browser has been closed")
                return {"average": "8.00"}

            service._ensure_enrich_browser = lambda: object()
            service._reset_enrich_browser = lambda: None
            service._open_competitor = fake_open
            service._read_1688_from_row = fake_source
            event = threading.Event()
            service._run_enrichment("B07LCLQ47D", "https://example.com/img.jpg", event)
            self.assertEqual(calls["open"], 2)
            self.assertEqual(calls["source"], 2)
            self.assertEqual(service.state["products"]["B07LCLQ47D"]["source_price"]["average"], "8.00")

    def test_closed_source_error_is_not_recorded_as_missing_price(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            service = SellerAutoService(DelaySettings(), output_dir=directory)
            with mock.patch(
                "seller_auto.service.open_1688_from_row",
                side_effect=RuntimeError("Target page, context or browser has been closed"),
            ):
                with self.assertRaises(RuntimeError):
                    service._read_1688_from_row(object(), object(), "B07LCLQ47D")


if __name__ == "__main__":
    unittest.main()
