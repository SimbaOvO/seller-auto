"""往 test.xlsx「第一批产品」后面追加一行，不另存新表。"""

from __future__ import annotations

import os
import re
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from seller_auto.sales import sales_count
from seller_auto.storage import public_url

TEMPLATE_SHEET = "第一批产品"
EXCHANGE_RATE = 0.004
CHANNEL_RATE = 0.15
MKD_LINK_COUNT = 10
LISTING_ROW_HEIGHT = 82

_FORMULAS = {
    "volume": "={length}*{width}*{height}/1000000",
    "channel_fee": "={channel_rate}*{price_rmb}",
    "freight": "={volume}*3500",
    "last_mile": "=({fba}+{fbm}+{transfer})*{rate}",
    "loss": "=(({price_rmb}*0.65)-{last_mile})*0.09",
    "withhold": "={price_rmb}*0.19",
    "price_rmb": "={price}*{rate}",
    "total_cost": "=({cost}+{channel_fee}+{freight}+{last_mile}+{warehouse}+{tax}+{loss}+{withhold})",
    "profit": "={price_rmb}-{total_cost}",
    "margin": "={profit}/{price_rmb}",
    "roi": "={profit}/({cost}+{freight})",
    "expected_volume": "={expected_qty}*{volume}",
    "expected_value": "={expected_qty}*{cost}",
}
_CONSTANTS = {
    "rate": EXCHANGE_RATE,
    "channel_rate": CHANNEL_RATE,
    "fbm": 0,
    "transfer": 0,
    "warehouse": 0,
    "tax": 0,
}


def template_path() -> Path:
    return Path(__file__).resolve().parent.parent / "test.xlsx"


def weight_kg(package: dict[str, Any] | None) -> float | None:
    if not package or package.get("weight") in (None, ""):
        return None
    value = Decimal(str(package["weight"]))
    unit = str(package.get("weight_unit") or "g").lower()
    if unit == "kg":
        kilograms = value
    elif unit == "g":
        kilograms = value / Decimal(1000)
    elif unit == "lb":
        kilograms = value * Decimal("0.45359237")
    elif unit == "oz":
        kilograms = value * Decimal("0.028349523125")
    else:
        return None
    return float(kilograms.quantize(Decimal("0.001"), rounding=ROUND_HALF_UP))


def dimension_cm(package: dict[str, Any] | None, key: str) -> float | None:
    if not package or package.get(key) in (None, ""):
        return None
    value = Decimal(str(package[key]))
    unit = str(package.get("unit") or "cm").lower()
    if unit == "in":
        value *= Decimal("2.54")
    elif unit != "cm":
        return None
    return float(value.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


def purchase_cost(product: dict[str, Any]) -> float | None:
    source = product.get("source_price") or {}
    average = source.get("average")
    if average in (None, ""):
        return None
    return float(Decimal(str(average)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def ranked_matches(product: dict[str, Any], limit: int = MKD_LINK_COUNT) -> list[dict[str, Any]]:
    """对比记录按销量倒序，取前 10。"""

    rows = [item for item in product.get("matches") or [] if item.get("url")]
    rows.sort(
        key=lambda item: (
            -sales_count(str(item.get("sales_text") or "")),
            -float(item.get("confidence") or item.get("similarity") or 0),
            str(item.get("candidate_id") or item.get("url") or ""),
        )
    )
    return rows[:limit]


def sales_cutoff(product: dict[str, Any], limit: int = MKD_LINK_COUNT) -> int | None:
    """同款销量已满 limit 条时，返回第 limit 名的销量；未满则没有门槛。"""

    rows = [item for item in product.get("matches") or [] if item.get("url")]
    if len(rows) < limit:
        return None
    top = ranked_matches(product, limit)
    if len(top) < limit:
        return None
    return sales_count(str(top[-1].get("sales_text") or ""))


def can_enter_top(product: dict[str, Any], sales_text: str, limit: int = MKD_LINK_COUNT) -> bool:
    """已记录同款是否挤得进销量前 10。写表取前 10 时用。"""

    cutoff = sales_cutoff(product, limit)
    if cutoff is None:
        return True
    return sales_count(sales_text) >= cutoff


def order_count(product: dict[str, Any]) -> int | None:
    """名称里「xx单」的数量，同时写入预计数量。"""

    orders = product.get("market_orders")
    if orders in (None, ""):
        return None
    try:
        count = int(orders)
    except (TypeError, ValueError):
        return None
    if count < 0:
        return None
    return count


def display_name(product: dict[str, Any]) -> str:
    """名称列。蓝鲸查到 30 天销量后，单量写在同一格的下一行。"""

    name = (product.get("chinese_name") or "").strip()
    count = order_count(product)
    if count is None:
        return name
    return f"{name}\n（{count}单）"


def fba_amount(raw: Any) -> int | float | None:
    if raw in (None, ""):
        return None
    try:
        number = Decimal(str(raw))
    except Exception:
        return None
    if number < 0:
        return None
    integral = number.to_integral_value()
    if number == integral:
        return int(integral)
    return float(number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def lowest_selling_price(matches: list[dict[str, Any]]) -> float | None:
    prices = []
    for item in matches:
        raw = item.get("price")
        if raw in (None, ""):
            continue
        try:
            prices.append(Decimal(str(raw)))
        except Exception:
            continue
    if not prices:
        return None
    return float(min(prices).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _name_key(value: Any) -> str:
    return str(value or "").split("\n", 1)[0].strip()


def _find_product_row(sheet: Worksheet, roles: dict[str, int], product: dict[str, Any]) -> int | None:
    sku = (product.get("keywords") or "").strip()
    name = (product.get("chinese_name") or "").strip()
    sku_column = roles.get("sku")
    name_column = roles.get("name")
    last = _next_data_row(sheet)
    hits: list[int] = []
    for row in range(2, last):
        cell_sku = str(sheet.cell(row, sku_column).value or "").strip() if sku_column else ""
        cell_name = _name_key(sheet.cell(row, name_column).value) if name_column else ""
        if sku and cell_sku == sku:
            hits.append(row)
        elif name and cell_name == name:
            hits.append(row)
    return hits[-1] if hits else None


def _fill_package_cells(
    sheet: Worksheet,
    row: int,
    roles: dict[str, int],
    product: dict[str, Any],
    *,
    replace_fba: bool = False,
) -> bool:
    package = product.get("package_size") or {}
    values: dict[str, Any] = {
        "weight": weight_kg(package),
        "length": dimension_cm(package, "length"),
        "width": dimension_cm(package, "width"),
        "height": dimension_cm(package, "height"),
        "cost": purchase_cost(product),
        "fba": fba_amount(product.get("fba_fee")),
        "expected_qty": order_count(product),
    }
    changed = False
    for role, value in values.items():
        column = roles.get(role)
        if not column or value in (None, ""):
            continue
        current = sheet.cell(row, column).value
        if current in (None, ""):
            sheet.cell(row, column, value)
            changed = True
        elif replace_fba and role == "fba" and current != value:
            sheet.cell(row, column, value)
            changed = True
    name_column = roles.get("name")
    if name_column and order_count(product) is not None:
        wanted = display_name(product)
        current = sheet.cell(row, name_column).value
        current_text = str(current or "")
        if wanted and wanted != current and "单" not in current_text:
            sheet.cell(row, name_column, wanted)
            changed = True
    return changed


def fill_package_rows(path: Path, products: list[dict[str, Any]], *, replace_fba: bool = False) -> int:
    """按 SKU 回填已有行的毛重、长宽高、采购成本和 FBA，不另加行。"""

    rows = [product for product in products if (product.get("keywords") or product.get("chinese_name"))]
    if not rows or not path.is_file():
        return 0
    workbook = load_workbook(path)
    sheet = workbook[TEMPLATE_SHEET]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    roles = _roles(headers)
    updated = 0
    for product in rows:
        row = _find_product_row(sheet, roles, product)
        if row is None:
            continue
        if _fill_package_cells(sheet, row, roles, product, replace_fba=replace_fba):
            updated += 1
    if updated:
        if not os.access(path, os.W_OK):
            path.chmod(path.stat().st_mode | 0o200)
        workbook.save(path)
    return updated


def write_product_workbook(path: Path, products: list[dict[str, Any]], template: Path | None = None) -> int:
    rows = [product for product in products if ranked_matches(product, limit=1)]
    if not rows:
        return 0
    source = path if path.is_file() else (template or template_path())
    workbook = load_workbook(source)
    sheet = workbook[TEMPLATE_SHEET]
    mkd_columns = _ensure_mkd_columns(sheet)
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    roles = _roles(headers)
    row = _next_data_row(sheet)
    for product in rows:
        _write_row(sheet, row, roles, mkd_columns, product)
        row += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_file() and not os.access(path, os.W_OK):
        path.chmod(path.stat().st_mode | 0o200)
    workbook.save(path)
    return len(rows)


def _headers(template: Path) -> list[Any]:
    workbook = load_workbook(template, read_only=True, data_only=False)
    try:
        sheet = workbook[TEMPLATE_SHEET]
        row = next(sheet.iter_rows(min_row=1, max_row=1))
        return [cell.value for cell in row]
    finally:
        workbook.close()


def _roles(headers: list[Any]) -> dict[str, int]:
    roles: dict[str, int] = {}
    for index, header in enumerate(headers, start=1):
        role = _role(header)
        if role and role not in roles:
            roles[role] = index
    return roles


def _mkd_columns(headers: list[Any]) -> list[int]:
    return [index for index, header in enumerate(headers, start=1) if "".join(str(header or "").split()) == "MKD链接"]


def _ensure_mkd_columns(sheet: Worksheet, count: int = MKD_LINK_COUNT) -> list[int]:
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    columns = _mkd_columns(headers)
    if len(columns) >= count:
        return columns[:count]
    insert_at = columns[-1] + 1 if columns else len(headers) + 1
    missing = count - len(columns)
    sheet.insert_cols(insert_at, amount=missing)
    _shift_formula_columns(sheet, insert_at, missing)
    sample_width = sheet.column_dimensions[get_column_letter(columns[0])].width if columns else 13
    for offset in range(missing):
        column = insert_at + offset
        sheet.cell(1, column, "MKD链接")
        sheet.column_dimensions[get_column_letter(column)].width = sample_width or 13
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    return _mkd_columns(headers)[:count]


_CELL_REF = re.compile(r"(\$?)([A-Z]+)(\$?)(\d+)")


def _shift_formula_columns(sheet: Worksheet, start_column: int, amount: int) -> None:
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=sheet.max_column):
        for cell in row:
            value = cell.value
            if not (isinstance(value, str) and value.startswith("=")):
                continue

            def replace(match: Any, start: int = start_column, shift: int = amount) -> str:
                dollar_col, letters, dollar_row, digits = match.groups()
                index = column_index_from_string(letters)
                if index >= start:
                    letters = get_column_letter(index + shift)
                return f"{dollar_col}{letters}{dollar_row}{digits}"

            cell.value = _CELL_REF.sub(replace, value)


def _retarget_formula(formula: str, source_row: int, target_row: int) -> str:
    """上一行的公式挪到新行，只改指向本行的单元格，0.23 这类常数不动。"""

    def replace(match: Any, source: int = source_row, target: int = target_row) -> str:
        dollar_col, letters, dollar_row, digits = match.groups()
        if not dollar_row and int(digits) == source:
            digits = str(target)
        return f"{dollar_col}{letters}{dollar_row}{digits}"

    return _CELL_REF.sub(replace, formula)


def _copied_formula(sheet: Worksheet, source_row: int | None, target_row: int, column: int) -> str:
    if not source_row or source_row < 2:
        return ""
    value = sheet.cell(source_row, column).value
    if not (isinstance(value, str) and value.startswith("=")):
        return ""
    return _retarget_formula(value, source_row, target_row)


def _next_data_row(sheet: Worksheet) -> int:
    row = 2
    while sheet.cell(row, 1).value not in (None, ""):
        row += 1
    return row


def _role(header: Any) -> str:
    key = "".join(str(header or "").split())
    if key.startswith("名称"):
        return "name"
    if key == "SKU":
        return "sku"
    if key.startswith("上架"):
        return "listing"
    if key.startswith("毛重"):
        return "weight"
    if key.startswith("长"):
        return "length"
    if key.startswith("宽"):
        return "width"
    if key.startswith("高"):
        return "height"
    if key.startswith("体积"):
        return "volume"
    if key.startswith("汇率"):
        return "rate"
    if key.startswith("售价"):
        return "price"
    if key.startswith("渠道比例"):
        return "channel_rate"
    if key.startswith("采购成本"):
        return "cost"
    if key.startswith("渠道费用"):
        return "channel_fee"
    if "运输方式" in key:
        return "shipping"
    if key.startswith("头程运费"):
        return "freight"
    if key.startswith("尾程"):
        return "last_mile"
    if key.startswith("FBA"):
        return "fba"
    if key.startswith("FBM"):
        return "fbm"
    if key.startswith("调拨费"):
        return "transfer"
    if key.startswith("仓库处理费"):
        return "warehouse"
    if key.startswith("税"):
        return "tax"
    if key.startswith("提现汇损"):
        return "loss"
    if key.startswith("代扣代缴"):
        return "withhold"
    if key.startswith("销售价格"):
        return "price_rmb"
    if key.startswith("成本合计"):
        return "total_cost"
    if key.startswith("毛利润"):
        return "profit"
    if key.startswith("毛利率"):
        return "margin"
    if key.startswith("投产比"):
        return "roi"
    if key.startswith("预计数量"):
        return "expected_qty"
    if key.startswith("预计方数"):
        return "expected_volume"
    if key.startswith("预计货值"):
        return "expected_value"
    return ""


def _write_row(
    sheet: Worksheet,
    row: int,
    roles: dict[str, int],
    mkd_columns: list[int],
    product: dict[str, Any],
) -> None:
    package = product.get("package_size") or {}
    top = ranked_matches(product)
    values: dict[str, Any] = {
        "name": display_name(product),
        "sku": (product.get("keywords") or "").strip(),
        "weight": weight_kg(package),
        "length": dimension_cm(package, "length"),
        "width": dimension_cm(package, "width"),
        "height": dimension_cm(package, "height"),
        "cost": purchase_cost(product),
        "price": lowest_selling_price(top),
        "shipping": "AR普货海运",
        "expected_qty": order_count(product),
    }
    values.update(_CONSTANTS)
    addresses = {role: f"{get_column_letter(column)}{row}" for role, column in roles.items()}
    source_row = row - 1 if row > 2 else None
    for role, column in roles.items():
        # 代扣代缴后续统一 19%，不沿用表里改过的 0.23。
        if role != "withhold":
            copied = _copied_formula(sheet, source_row, row, column)
            if copied:
                sheet.cell(row, column, copied)
                continue
        if role == "listing":
            continue
        if role == "fba":
            fee = fba_amount(product.get("fba_fee"))
            if fee is not None:
                sheet.cell(row, column, fee)
            continue
        if role in _FORMULAS:
            formula = _formula(role, addresses)
            if formula:
                sheet.cell(row, column, formula)
            continue
        value = values.get(role)
        if value in (None, ""):
            continue
        sheet.cell(row, column, value)
    _style_filled_columns(sheet, source_row, row, roles)
    for column, match in zip(mkd_columns, top):
        url = public_url(str(match.get("url") or ""), str(match.get("key") or match.get("candidate_id") or ""))
        if url:
            sheet.cell(row, column, url)
    _center_row(sheet, row, roles, mkd_columns)
    _embed_listing_image(sheet, row, roles.get("listing"), product.get("image_path") or "")
    sheet.row_dimensions[row].height = LISTING_ROW_HEIGHT


def _center_row(sheet: Worksheet, row: int, roles: dict[str, int], mkd_columns: list[int]) -> None:
    name_column = roles.get("name")
    for column in list(roles.values()) + list(mkd_columns):
        cell = sheet.cell(row, column)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=column == name_column,
        )


def _earlier_fill(sheet: Worksheet, row: int, column: int):
    for source in range(row - 1, 1, -1):
        fill = sheet.cell(source, column).fill
        if fill and fill.patternType:
            return fill
    return None


def _style_filled_columns(sheet: Worksheet, source_row: int | None, row: int, roles: dict[str, int]) -> None:
    """采购成本和 FBA 运费沿用表格里已有的底色。投产比不填色。"""

    _ = source_row
    for role in ("cost", "fba"):
        column = roles.get(role)
        if not column:
            continue
        fill = _earlier_fill(sheet, row, column)
        if fill is not None:
            sheet.cell(row, column).fill = copy(fill)
    roi = roles.get("roi")
    if roi:
        cell = sheet.cell(row, roi)
        cell.fill = PatternFill(fill_type=None)
        cell.number_format = "0.00%"


def _embed_listing_image(sheet: Worksheet, row: int, column: int | None, image_path: str) -> None:
    if not column or not image_path:
        return
    path = Path(image_path)
    if not path.is_file():
        return
    image = ExcelImage(str(path))
    image.width = 78
    image.height = 70
    image.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=column - 1, colOff=0, row=row - 1, rowOff=0),
        to=AnchorMarker(col=column, colOff=0, row=row, rowOff=0),
    )
    sheet.add_image(image)


def _formula(role: str, addresses: dict[str, str]) -> str:
    try:
        return _FORMULAS[role].format(**addresses)
    except KeyError:
        return ""
